#!/usr/bin/env python3
"""
Scraper de veille concurrentielle - Alliance Nautique 66
Client : Marine Center
Auteur : Équipe développement
"""

import argparse
import json
import logging
import re
import time
from datetime import date
from pathlib import Path
from urllib.parse import urljoin, urlparse

from bs4 import BeautifulSoup
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────────────────────────────────────
# Configuration
# ─────────────────────────────────────────────────────────────────────────────

BASE_URL = "https://www.an66.fr"
NEUFS_URL = "https://www.an66.fr/bateaux-neufs/"
OCCASIONS_URL = "https://www.an66.fr/occasions/"
DATA_FILE = Path("data/boats_history.json")
OUTPUT_DIR = Path("output")
DEBUG_DIR = Path("debug")
TODAY = date.today().isoformat()

# ─────────────────────────────────────────────────────────────────────────────
# Logging
# ─────────────────────────────────────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger(__name__)

# ─────────────────────────────────────────────────────────────────────────────
# Data persistence
# ─────────────────────────────────────────────────────────────────────────────


def load_history() -> dict:
    if DATA_FILE.exists():
        with open(DATA_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {"neufs": {}, "occasions": {}}


def save_history(history: dict) -> None:
    DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(history, f, ensure_ascii=False, indent=2)
    logger.info(f"Historique sauvegardé : {DATA_FILE}")


# ─────────────────────────────────────────────────────────────────────────────
# Utilitaires prix
# ─────────────────────────────────────────────────────────────────────────────


def parse_price(text: str) -> float | None:
    """
    Parse un prix au format français : '495 000,00 €' → 495000.0
    Espace = séparateur de milliers, virgule = séparateur décimal.
    """
    if not text:
        return None
    # Supprime € et espaces
    clean = re.sub(r"[€\s]", "", text).strip()
    if not clean:
        return None
    if "," in clean:
        parts = clean.split(",")
        integer_part = re.sub(r"[^\d]", "", parts[0])
        # La virgule est décimale seulement si 1-2 chiffres après
        if len(parts) > 1 and len(re.sub(r"[^\d]", "", parts[1])) <= 2:
            decimal_part = re.sub(r"[^\d]", "", parts[1])
            full = integer_part + "." + decimal_part
        else:
            full = integer_part
    else:
        full = re.sub(r"[^\d]", "", clean)
    return float(full) if full else None


def format_price(price: float | None) -> str:
    if price is None:
        return "N/C"
    return f"{price:,.0f} €".replace(",", " ")


# ─────────────────────────────────────────────────────────────────────────────
# Scraping web
# ─────────────────────────────────────────────────────────────────────────────

# Sélecteurs CSS par ordre de priorité pour les cartes bateau
CARD_SELECTORS = [
    # Patterns spécifiques aux sites de concessionnaires nautiques
    ".annonce-item", ".annonce", ".boat-card", ".bateau-item", ".bateau",
    ".listing-item", ".listing-boat", ".fiche-bateau",
    # WordPress classique
    "article.post", "article.product",
    # Patterns génériques avec "annonce" ou "boat" dans la classe
    '[class*="annonce"]', '[class*="bateau"]', '[class*="boat"]',
    # WooCommerce / shop
    ".products li.product", ".product-item",
    # Grilles génériques
    ".grid-item", ".col-item",
    "article", ".col-md-6", ".col-sm-6", ".col-lg-4", ".col-md-4",
]

PRICE_SELECTORS = [
    ".prix", ".price", "#prix", '[class*="prix"]', '[class*="price"]',
    ".tarif", '[class*="tarif"]', ".montant", '[class*="montant"]',
    "strong", ".woocommerce-Price-amount",
]

TITLE_SELECTORS = [
    "h1", "h2", "h3", "h4",
    ".titre", ".title", ".nom", ".modele",
    '[class*="titre"]', '[class*="title"]', '[class*="modele"]',
]

MOTOR_BRANDS = (
    r"Mercury|Yamaha|Honda|Volvo|Mercruiser|Suzuki|Evinrude|Johnson|"
    r"Tohatsu|Yanmar|Nanni|Perkins|Caterpillar|Cummins|Torqeedo"
)


def get_page_html(page, url: str, debug: bool = False) -> str | None:
    """Charge une URL avec Playwright et retourne le HTML."""
    for attempt in range(3):
        try:
            page.goto(url, wait_until="load", timeout=60000)
            # Attend que le contenu principal soit visible
            time.sleep(3 + attempt * 2)
            html = page.content()
            if debug:
                DEBUG_DIR.mkdir(parents=True, exist_ok=True)
                slug = re.sub(r"[^\w]", "_", urlparse(url).path)[:60]
                dbg_file = DEBUG_DIR / f"{slug}.html"
                dbg_file.write_text(html, encoding="utf-8")
                logger.debug(f"HTML sauvegardé : {dbg_file}")
            return html
        except PlaywrightTimeoutError:
            logger.warning(f"Timeout (tentative {attempt+1}/3) pour {url}")
            if attempt < 2:
                try:
                    page.goto(url, wait_until="domcontentloaded", timeout=60000)
                    time.sleep(4 + attempt * 2)
                except Exception:
                    pass
        except Exception as e:
            logger.error(f"Erreur chargement {url} : {e}")
            return None
    return None


def extract_boat_from_card(element, base_url: str) -> dict | None:
    """Extrait les infos basiques d'une carte bateau (listing)."""
    link = element.find("a", href=True)
    if not link:
        return None

    href = link["href"]
    if any(skip in href for skip in ["#", "javascript:", "mailto:", "tel:"]):
        return None

    url = urljoin(base_url, href)
    if url.rstrip("/") == base_url.rstrip("/"):
        return None

    boat: dict = {"url": url}

    # Titre / Modèle
    for sel in TITLE_SELECTORS:
        el = element.select_one(sel)
        if el:
            text = el.get_text(strip=True)
            if text and len(text) > 2:
                boat["modele"] = text
                break
    if not boat.get("modele"):
        link_text = link.get_text(strip=True)
        if link_text and len(link_text) > 2:
            boat["modele"] = link_text

    # Prix depuis la carte
    for sel in PRICE_SELECTORS:
        el = element.select_one(sel)
        if el:
            text = el.get_text(strip=True)
            if "€" in text or re.search(r"\d{4,}", text):
                boat["prix_text"] = text
                boat["prix"] = parse_price(text)
                break

    return boat if boat.get("modele") or boat.get("prix") else None


def find_boat_links(soup: BeautifulSoup, base_url: str) -> list[dict]:
    """Trouve tous les liens de bateaux sur une page de listing."""
    seen: set[str] = set()
    boats: list[dict] = []

    # Stratégie 0 : détection via le bouton "Plus d'informations" (spécifique AN66)
    info_buttons = soup.find_all("a", string=re.compile(r"plus d.information", re.I))
    if not info_buttons:
        info_buttons = [
            a for a in soup.find_all("a", href=True)
            if re.search(r"plus d.information|détail|voir la fiche", a.get_text(), re.I)
        ]
    if info_buttons:
        logger.info(f"Stratégie bouton 'Plus d'informations' : {len(info_buttons)} bateaux")
        for btn in info_buttons:
            url = urljoin(base_url, btn["href"])
            if url in seen:
                continue
            seen.add(url)
            # Remonte dans le DOM pour trouver le conteneur du bateau
            container = btn.parent
            for _ in range(6):
                if container is None:
                    break
                title_el = container.find(["h1", "h2", "h3", "h4"])
                if title_el and title_el.get_text(strip=True):
                    break
                container = container.parent
            boat: dict = {"url": url}
            if container:
                title_el = container.find(["h1", "h2", "h3", "h4"])
                if title_el:
                    boat["modele"] = title_el.get_text(strip=True)
                text = container.get_text(separator=" ")
                prix_m = re.search(r"Prix\s*[:\s]*([0-9][0-9\s\.,]+\s*€)", text, re.I)
                if prix_m:
                    boat["prix_text"] = prix_m.group(1).strip()
                    boat["prix"] = parse_price(prix_m.group(1))
                annee_m = re.search(r"Ann[ée]e\s*[:\s]*(\d{4})", text, re.I)
                if annee_m:
                    boat["annee"] = int(annee_m.group(1))
            boats.append(boat)
        if boats:
            return boats

    # Stratégie 1 : cartes structurées
    for selector in CARD_SELECTORS:
        elements = soup.select(selector)
        if len(elements) < 2:
            continue
        with_links = [el for el in elements if el.find("a", href=True)]
        if not with_links or len(with_links) / len(elements) < 0.5:
            continue
        logger.info(f"Sélecteur retenu : '{selector}' → {len(with_links)} éléments")
        for el in with_links:
            boat = extract_boat_from_card(el, base_url)
            if boat and boat["url"] not in seen:
                seen.add(boat["url"])
                boats.append(boat)
        if boats:
            return boats

    # Stratégie 2 : extraction par liens ressemblant à des fiches bateaux
    logger.warning("Aucun sélecteur de carte valide — repli sur extraction de liens")
    keywords = ["annonce", "bateau", "occasion", "voilier", "catamaran",
                "moteur", "semi-rigide", "pneumatique", "id_annonce", "id="]
    for a in soup.find_all("a", href=True):
        href = a["href"]
        url = urljoin(base_url, href)
        if url in seen:
            continue
        if any(kw in href.lower() for kw in keywords):
            seen.add(url)
            text = a.get_text(strip=True)
            if text and len(text) > 3:
                boats.append({"url": url, "modele": text})

    return boats


def find_next_page(soup: BeautifulSoup, current_url: str, base_url: str) -> str | None:
    """Retourne l'URL de la page suivante (pagination) ou None."""
    patterns = [
        'a[rel="next"]', ".next a", ".suivant a",
        '[class*="next"] a', '[class*="suivant"] a',
        'a[aria-label*="suivant"]', 'a[aria-label*="next"]',
        ".nav-links a.next", ".pagination a.next",
    ]
    for p in patterns:
        el = soup.select_one(p)
        if el and el.get("href"):
            return urljoin(base_url, el["href"])

    # Cherche le numéro de page courant et retourne le suivant
    page_links = soup.select(".pagination a, [class*='pag'] a, [class*='page'] a")
    for i, lnk in enumerate(page_links):
        classes = " ".join(lnk.get("class", []))
        if "current" in classes or "active" in classes:
            if i + 1 < len(page_links):
                href = page_links[i + 1].get("href")
                if href:
                    return urljoin(base_url, href)
    return None


def scrape_detail_page(page, url: str, debug: bool = False) -> dict:
    """
    Visite la page détail d'un bateau AN66 et extrait :
    modele, prix, année, motorisation, heures moteur.
    Structure réelle du site : table.table-striped avec <th>Label</th><td>Valeur</td>
    """
    logger.info(f"  → {url}")
    html = get_page_html(page, url, debug=debug)
    if not html:
        return {}

    soup = BeautifulSoup(html, "html.parser")
    details: dict = {}

    # ── Modèle : premier <h3> qui n'est pas "Les photos..." ──────────────────
    for h3 in soup.find_all("h3"):
        text = h3.get_text(strip=True)
        if text and len(text) > 2 and "photo" not in text.lower():
            details["modele"] = text
            break

    # ── Table de specs AN66 : <table class="table table-striped"> ─────────────
    # Labels connus : Marque, Modèle, Moteur, Année, Prix, Heures, Heure moteur
    spec_table = soup.select_one("table.table-striped, table.table")
    if spec_table:
        for row in spec_table.find_all("tr"):
            th = row.find("th")
            td = row.find("td")
            if not th:
                continue
            label = th.get_text(strip=True).lower().strip()
            # Valeur : dans <td> si présent, sinon texte du <tr> sans le label
            if td:
                value = td.get_text(strip=True)
            else:
                value = row.get_text(strip=True).replace(th.get_text(strip=True), "", 1).strip()

            if not value:
                continue

            if label == "prix" and not details.get("prix"):
                details["prix_text"] = value
                details["prix"] = parse_price(value)
            elif label == "moteur" and not details.get("motorisation"):
                details["motorisation"] = value
            elif label in ("année", "annee") and not details.get("annee"):
                m = re.search(r"\d{4}", value)
                if m and 1980 <= int(m.group()) <= 2030:
                    details["annee"] = int(m.group())
            elif re.search(r"heure", label) and not details.get("heures_moteur"):
                m = re.search(r"\d+", value)
                if m:
                    details["heures_moteur"] = m.group()
            elif label == "modèle" and not details.get("modele"):
                details["modele"] = value

    # ── Fallback prix si table vide ───────────────────────────────────────────
    if not details.get("prix"):
        full_text = soup.get_text(separator=" ", strip=True)
        m = re.search(r"(\d[\d\s]{2,}[,.]?\d{0,2})\s*€", full_text)
        if m:
            details["prix_text"] = m.group().strip()
            details["prix"] = parse_price(m.group())

    return details


def _extract_from_specs(soup: BeautifulSoup, details: dict) -> None:
    """Non utilisé — gardé pour compatibilité."""
    pass


def scrape_category(
    page, url: str, label: str, debug: bool = False
) -> list[dict]:
    """
    Scrape toutes les pages d'une catégorie (avec pagination).
    Retourne la liste brute des bateaux trouvés.
    """
    all_boats: list[dict] = []
    current_url: str | None = url
    page_num = 1
    seen_urls: set[str] = set()

    while current_url:
        logger.info(f"[{label}] Page {page_num} : {current_url}")
        html = get_page_html(page, current_url, debug=debug)
        if not html:
            break

        soup = BeautifulSoup(html, "html.parser")
        boat_stubs = find_boat_links(soup, BASE_URL)

        if not boat_stubs:
            logger.warning(
                f"Aucun bateau trouvé page {page_num}. "
                "Lancez avec --debug pour inspecter le HTML."
            )
            break

        logger.info(f"  {len(boat_stubs)} bateaux trouvés sur la page {page_num}")

        for stub in boat_stubs:
            boat_url = stub["url"]
            if boat_url in seen_urls:
                continue
            seen_urls.add(boat_url)

            detail = scrape_detail_page(page, boat_url, debug=debug)
            boat = {**stub}
            # Le détail de la page prime sur le listing, sauf pour le modèle :
            # on garde le modèle du listing s'il existe, sinon on prend celui du détail
            for k, v in detail.items():
                if v is not None:
                    if k == "modele" and boat.get("modele"):
                        continue  # on garde le titre du listing
                    boat[k] = v
            boat.setdefault("modele", detail.get("modele", "N/C"))
            boat.setdefault("annee", None)
            boat.setdefault("motorisation", None)
            boat.setdefault("heures_moteur", None)
            boat.setdefault("prix", None)
            boat.setdefault("prix_text", None)
            all_boats.append(boat)
            time.sleep(1)  # politesse envers le serveur

        next_url = find_next_page(soup, current_url, BASE_URL)
        if next_url and next_url != current_url:
            current_url = next_url
            page_num += 1
            time.sleep(2)
        else:
            break

    return all_boats


# ─────────────────────────────────────────────────────────────────────────────
# Gestion de l'historique et calcul des baisses de prix
# ─────────────────────────────────────────────────────────────────────────────


def compute_price_drops(historique: list[dict]) -> list[dict]:
    """Calcule les baisses de prix depuis l'historique."""
    drops = []
    for i in range(1, len(historique)):
        prev = historique[i - 1]["prix"]
        curr = historique[i]["prix"]
        if prev and curr and curr < prev - 0.01:
            drops.append(
                {
                    "montant": prev - curr,
                    "pourcentage": ((prev - curr) / prev) * 100,
                    "prix_avant": prev,
                    "prix_apres": curr,
                    "date": historique[i]["date"],
                }
            )
    return drops


def update_history(
    history: dict, category: str, scraped_boats: list[dict]
) -> list[dict]:
    """
    Met à jour l'historique avec les bateaux scrapés.
    Retourne la liste enrichie avec date_parution et baisses_de_prix.
    """
    cat = history.setdefault(category, {})
    enriched: list[dict] = []

    for boat in scraped_boats:
        url = boat["url"]
        current_price = boat.get("prix")

        if url not in cat:
            # Nouveau bateau
            cat[url] = {
                "modele": boat.get("modele"),
                "annee": boat.get("annee"),
                "motorisation": boat.get("motorisation"),
                "heures_moteur": boat.get("heures_moteur"),
                "date_parution": TODAY,
                "historique_prix": (
                    [{"prix": current_price, "date": TODAY}]
                    if current_price
                    else []
                ),
            }
            logger.info(f"  + Nouveau : {boat.get('modele', url)}")
        else:
            existing = cat[url]
            # Toujours mettre à jour motorisation/heures/année avec la dernière valeur scrapée
            for field in ("annee", "motorisation", "heures_moteur"):
                if boat.get(field) is not None:
                    existing[field] = boat[field]
            # Pour le modèle on garde l'existant s'il est déjà renseigné
            if not existing.get("modele") and boat.get("modele"):
                existing["modele"] = boat["modele"]

            # Détecte un changement de prix
            hist = existing["historique_prix"]
            if current_price:
                if hist:
                    last_price = hist[-1]["prix"]
                    if last_price and abs(current_price - last_price) > 0.01:
                        hist.append({"prix": current_price, "date": TODAY})
                        direction = "↓" if current_price < last_price else "↑"
                        logger.info(
                            f"  {direction} Prix modifié : {boat.get('modele')} "
                            f"{format_price(last_price)} → {format_price(current_price)}"
                        )
                else:
                    hist.append({"prix": current_price, "date": TODAY})

        rec = cat[url]
        enriched.append(
            {
                "url": url,
                "modele": rec.get("modele") or boat.get("modele", "N/C"),
                "annee": rec.get("annee") or boat.get("annee"),
                "motorisation": rec.get("motorisation") or boat.get("motorisation"),
                "heures_moteur": rec.get("heures_moteur") or boat.get("heures_moteur"),
                "prix_actuel": current_price,
                "prix_text": boat.get("prix_text"),
                "date_parution": rec["date_parution"],
                "historique_prix": rec["historique_prix"],
                "baisses_de_prix": compute_price_drops(rec["historique_prix"]),
            }
        )

    return enriched


# ─────────────────────────────────────────────────────────────────────────────
# Génération du rapport Excel
# ─────────────────────────────────────────────────────────────────────────────

# Palette de couleurs Marine Center
C_BLUE_DARK = "1A3D6E"    # Bleu marine — en-tête neufs
C_GOLD_DARK = "8B5E14"    # Or foncé — en-tête occasions
C_WHITE = "FFFFFF"
C_BLUE_LIGHT = "DCE9F7"   # Fond lignes paires neufs
C_GOLD_LIGHT = "FFF4DC"   # Fond lignes paires occasions
C_ROW_WHITE = "FFFFFF"
C_PRICE_DROP = "FFD966"   # Jaune or — baisse de prix
C_NEW_ENTRY = "D9EAD3"    # Vert clair — première apparition


COLUMNS: list[tuple[str, int]] = [
    ("Modèle / Nom", 32),
    ("Année", 8),
    ("Motorisation", 28),
    ("Heures Moteur", 14),
    ("Prix Actuel", 14),
    ("Date Parution", 13),
    ("Baisse(s) de Prix", 20),
    ("Date de Baisse", 14),
    ("% Baisse", 10),
    ("URL", 50),
]


def _thin_border() -> Border:
    thin = Side(style="thin")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _header_fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _create_sheet(
    wb: openpyxl.Workbook,
    sheet_name: str,
    boats: list[dict],
    is_neuf: bool,
) -> None:
    ws = wb.create_sheet(title=sheet_name)
    h_color = C_BLUE_DARK if is_neuf else C_GOLD_DARK
    row_alt = C_BLUE_LIGHT if is_neuf else C_GOLD_LIGHT
    border = _thin_border()
    ncols = len(COLUMNS)

    # ── Ligne titre ──────────────────────────────────────────────────────────
    ws.merge_cells(f"A1:{get_column_letter(ncols)}1")
    title_cell = ws["A1"]
    title_cell.value = (
        f"{'Bateaux Neufs' if is_neuf else 'Occasions'} · "
        f"Alliance Nautique 66  |  Veille Marine Center  |  {TODAY}"
    )
    title_cell.font = Font(bold=True, size=13, color=C_WHITE)
    title_cell.fill = _header_fill(h_color)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # ── Ligne en-têtes colonnes ───────────────────────────────────────────────
    for ci, (col_name, col_width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=2, column=ci, value=col_name)
        cell.font = Font(bold=True, color=C_WHITE, size=10)
        cell.fill = _header_fill(h_color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(ci)].width = col_width
    ws.row_dimensions[2].height = 36

    if not boats:
        ws.cell(row=3, column=1, value="Aucun bateau trouvé.").font = Font(
            italic=True, color="888888"
        )
        return

    # ── Lignes données ────────────────────────────────────────────────────────
    for ri, boat in enumerate(boats):
        row = 3 + ri
        is_odd = ri % 2 == 0
        baisses = boat.get("baisses_de_prix", [])
        has_drop = bool(baisses)

        if has_drop:
            fill_color = C_PRICE_DROP
        elif boat["date_parution"] == TODAY:
            fill_color = C_NEW_ENTRY
        else:
            fill_color = row_alt if is_odd else C_ROW_WHITE

        row_fill = PatternFill("solid", fgColor=fill_color)

        # Formatage des baisses
        if baisses:
            baisses_str = " | ".join(
                f"-{b['montant']:,.0f} €".replace(",", " ") for b in baisses
            )
            dates_str = " | ".join(b["date"] for b in baisses)
            pct_str = " | ".join(f"-{b['pourcentage']:.1f}%" for b in baisses)
        else:
            baisses_str = dates_str = pct_str = ""

        prix_display = (
            format_price(boat.get("prix_actuel"))
            if boat.get("prix_actuel")
            else (boat.get("prix_text") or "N/C")
        )

        row_values = [
            boat.get("modele", "N/C"),
            boat.get("annee", ""),
            boat.get("motorisation", ""),
            boat.get("heures_moteur", ""),
            prix_display,
            boat.get("date_parution", TODAY),
            baisses_str,
            dates_str,
            pct_str,
            boat.get("url", ""),
        ]

        for ci, val in enumerate(row_values, 1):
            cell = ws.cell(row=row, column=ci, value=val)
            cell.fill = row_fill
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=(ci in (3, 7, 10)),
            )

        # Hyperlien sur la colonne URL
        url_cell = ws.cell(row=row, column=10)
        if boat.get("url"):
            url_cell.hyperlink = boat["url"]
            url_cell.font = Font(color="0563C1", underline="single")

        ws.row_dimensions[row].height = 20

    # ── Gel des en-têtes ─────────────────────────────────────────────────────
    ws.freeze_panes = "A3"

    # ── Filtre automatique ────────────────────────────────────────────────────
    ws.auto_filter.ref = f"A2:{get_column_letter(ncols)}{2 + len(boats)}"

    # ── Légende ───────────────────────────────────────────────────────────────
    leg_row = 3 + len(boats) + 2
    ws.cell(row=leg_row, column=1, value="Légende :").font = Font(bold=True)

    for offset, (color, label) in enumerate(
        [
            (C_PRICE_DROP, "Baisse de prix détectée"),
            (C_NEW_ENTRY, "Nouvelle annonce (aujourd'hui)"),
            (row_alt, "Ligne normale"),
        ],
        1,
    ):
        cell = ws.cell(row=leg_row + offset, column=1, value=f"  {label}")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.border = _thin_border()


def create_excel_report(neufs: list[dict], occasions: list[dict]) -> Path:
    """Génère le fichier Excel de veille."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    filename = OUTPUT_DIR / f"veille_an66_{TODAY}.xlsx"

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _create_sheet(wb, "Bateaux Neufs", neufs, is_neuf=True)
    _create_sheet(wb, "Occasions", occasions, is_neuf=False)
    wb.save(filename)
    logger.info(f"Rapport Excel : {filename}")
    return filename


# ─────────────────────────────────────────────────────────────────────────────
# Point d'entrée
# ─────────────────────────────────────────────────────────────────────────────


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Scraper de veille - Alliance Nautique 66 pour Marine Center"
    )
    parser.add_argument(
        "--debug",
        action="store_true",
        help="Sauvegarde les HTML bruts dans debug/ et active les logs détaillés",
    )
    parser.add_argument(
        "--headless",
        action="store_true",
        default=True,
        help="Exécuter le navigateur en mode headless (par défaut)",
    )
    parser.add_argument(
        "--no-headless",
        dest="headless",
        action="store_false",
        help="Afficher la fenêtre du navigateur (utile pour debug)",
    )
    args = parser.parse_args()

    if args.debug:
        logger.setLevel(logging.DEBUG)

    logger.info("═" * 60)
    logger.info("  Scraper Alliance Nautique 66 — Marine Center")
    logger.info(f"  Date : {TODAY}")
    logger.info("═" * 60)

    history = load_history()

    with sync_playwright() as pw:
        browser = pw.chromium.launch(
            headless=args.headless,
            args=[
                "--no-sandbox",
                "--disable-blink-features=AutomationControlled",
                "--disable-dev-shm-usage",
                "--disable-web-security",
            ],
        )

        context = browser.new_context(
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/122.0.0.0 Safari/537.36"
            ),
            viewport={"width": 1920, "height": 1080},
            locale="fr-FR",
            timezone_id="Europe/Paris",
        )

        # Bloque images/fonts pour accélérer le scraping
        context.route(
            "**/*.{png,jpg,jpeg,gif,webp,svg,ico,woff,woff2,ttf,eot,mp4,mp3}",
            lambda route: route.abort(),
        )

        page = context.new_page()
        # Masque navigator.webdriver pour éviter la détection bot
        page.add_init_script(
            "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"
        )

        logger.info("\n[1/2] Bateaux Neufs")
        neufs_raw = scrape_category(page, NEUFS_URL, "Neufs", debug=args.debug)
        logger.info(f"→ {len(neufs_raw)} bateaux neufs trouvés")

        logger.info("\n[2/2] Occasions")
        occasions_raw = scrape_category(page, OCCASIONS_URL, "Occasions", debug=args.debug)
        logger.info(f"→ {len(occasions_raw)} occasions trouvées")

        context.close()
        browser.close()

    neufs_final = update_history(history, "neufs", neufs_raw)
    occasions_final = update_history(history, "occasions", occasions_raw)

    save_history(history)

    excel_file = create_excel_report(neufs_final, occasions_final)

    logger.info("\n" + "═" * 60)
    logger.info(f"  Rapport : {excel_file}")
    logger.info(f"  Bateaux neufs  : {len(neufs_final)}")
    logger.info(f"  Occasions      : {len(occasions_final)}")
    logger.info("═" * 60)


if __name__ == "__main__":
    main()
